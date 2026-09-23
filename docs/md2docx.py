# -*- coding: utf-8 -*-
"""md2docx.py — Konversi dokumen Markdown (docs/09-analisis-cv-rag.md) ke DOCX.

Mendukung: heading (#..####), paragraf, blockquote, tabel pipe, bullet (-) &
numbered (1.) list, fenced code block, gambar (![](path)), bold **x**, inline `code`.

Gambar disisipkan dari path relatif terhadap file .md.

Usage:
  apps\\ai\\.venv\\Scripts\\python.exe docs\\md2docx.py docs\\09-analisis-cv-rag.md
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

CODE_FONT = "Consolas"


def _add_runs(paragraph, text: str) -> None:
    """Parse inline **bold** dan `code` lalu tambahkan run."""
    # pecah token **bold** dan `code`
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        token = m.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        else:  # `code`
            r = paragraph.add_run(token[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run("\n".join(lines))
    r.font.name = CODE_FONT
    r.font.size = Pt(8.5)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cells[ci].text = ""
            para = cells[ci].paragraphs[0]
            _add_runs(para, text)
            if ri == 0:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


def _split_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.strip() for c in line.split("|")]


def convert(md_path: Path, out_path: Path) -> None:
    base = md_path.parent
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    # style default
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            block: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            _add_code_block(doc, block)
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=min(level, 4))
            i += 1
            continue

        # tabel pipe
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows: list[list[str]] = []
            header = _split_row(lines[i])
            rows.append(header)
            i += 2  # lewati separator
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            continue

        # blockquote
        if stripped.startswith(">"):
            q: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                q.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _add_runs(p, " ".join(x for x in q if x.strip()))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            continue

        # gambar
        m = re.match(r"^!\[.*?\]\((.*?)\)\s*$", stripped)
        if m:
            img_path = (base / m.group(1)).resolve()
            if img_path.exists():
                doc.add_picture(str(img_path), width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[gambar tidak ditemukan: {m.group(1)}]")
            i += 1
            continue

        # horizontal rule
        if stripped in ("---", "***", "___"):
            doc.add_paragraph()
            i += 1
            continue

        # bullet list
        if re.match(r"^[-*]\s+", stripped):
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                text = re.sub(r"^[-*]\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Bullet")
                _add_runs(p, text)
                i += 1
            continue

        # numbered list
        if re.match(r"^\d+\.\s+", stripped):
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                text = re.sub(r"^\d+\.\s+", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                _add_runs(p, text)
                i += 1
            continue

        # paragraf kosong
        if not stripped:
            i += 1
            continue

        # paragraf biasa
        p = doc.add_paragraph()
        _add_runs(p, stripped)
        i += 1

    doc.save(str(out_path))
    print(f"OK -> {out_path}")


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("docs/09-analisis-cv-rag.md")
    dst = src.with_suffix(".docx")
    convert(src.resolve(), dst.resolve())
